"""Belge okuma, parçalama (chunking) ve indeksleme."""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional

import numpy as np

from . import store
from .config import DOCS_DIR
from .embeddings import BaseEmbedder

SUPPORTED = {".txt", ".md", ".markdown", ".pdf", ".docx", ".csv", ".json", ".log", ".rst"}


def read_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    for encoding in ("utf-8", "utf-8-sig", "cp1254", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_bytes().decode("utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("PDF okumak için 'pypdf' paketi gerekli.") from exc
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Sayfa {index}]\n{text}")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("DOCX okumak için 'python-docx' paketi gerekli.") from exc
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


HEADING_RE = re.compile(r"^(#{1,6}\s+.+|\[Sayfa \d+\])\s*$", re.MULTILINE)


def split_sections(text: str) -> List[tuple[str, str]]:
    """Metni başlıklarına göre (başlık, gövde) bölümlerine ayırır.

    Başlıklı belgelerde her bölüm ayrı parçalanır; böylece farklı konular tek
    bir vektörde birbirine karışmaz ve arama isabeti belirgin biçimde artar.
    """
    matches = list(HEADING_RE.finditer(text))
    if not matches:
        return [("", text)]

    sections: List[tuple[str, str]] = []
    if matches[0].start() > 0:
        intro = text[: matches[0].start()].strip()
        if intro:
            sections.append(("", intro))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        heading = match.group(0).lstrip("#").strip()
        body = text[match.end():end].strip()
        if body or heading:
            sections.append((heading, body))
    return sections


def chunk_text(text: str, size: int = 900, overlap: int = 150) -> List[str]:
    """Metni başlık ve paragraf sınırlarına saygı göstererek parçalara böler."""
    text = clean_text(text)
    if not text:
        return []

    chunks: List[str] = []
    for heading, body in split_sections(text):
        prefix = f"{heading}\n\n" if heading else ""
        for piece in _chunk_block(body, size=max(120, size - len(prefix)), overlap=overlap):
            chunks.append(f"{prefix}{piece}".strip())
    return [c for c in chunks if c.strip()]


def _chunk_block(text: str, size: int, overlap: int) -> List[str]:
    """Tek bir bölümü örtüşmeli parçalara böler."""
    text = text.strip()
    if not text:
        return []

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    units: List[str] = []
    for paragraph in paragraphs:
        if len(paragraph) <= size:
            units.append(paragraph)
            continue
        # çok uzun paragrafı cümlelere böl
        sentences = re.split(r"(?<=[.!?…])\s+", paragraph)
        buffer = ""
        for sentence in sentences:
            if len(buffer) + len(sentence) + 1 <= size:
                buffer = f"{buffer} {sentence}".strip()
            else:
                if buffer:
                    units.append(buffer)
                while len(sentence) > size:  # tek cümle bile uzunsa kes
                    units.append(sentence[:size])
                    sentence = sentence[size:]
                buffer = sentence
        if buffer:
            units.append(buffer)

    chunks: List[str] = []
    current = ""
    for unit in units:
        candidate = f"{current}\n\n{unit}".strip() if current else unit
        if len(candidate) <= size or not current:
            current = candidate
        else:
            chunks.append(current)
            tail = current[-overlap:] if overlap > 0 else ""
            if tail:
                cut = tail.find(" ")
                tail = tail[cut + 1:] if cut != -1 else tail
            current = f"{tail}\n\n{unit}".strip() if tail else unit
    if current:
        chunks.append(current)
    return [c for c in chunks if c.strip()]


def index_text(name: str, text: str, source: str, kind: str,
               embedder: BaseEmbedder, settings: Dict[str, Any]) -> Dict[str, Any]:
    chunks = chunk_text(
        text,
        size=int(settings.get("chunk_size", 900)),
        overlap=int(settings.get("chunk_overlap", 150)),
    )
    if not chunks:
        raise ValueError(f"'{name}' içinden metin çıkarılamadı.")
    # Belge adı da gömmeye katılır: "alaka-kapisi.md" içindeki parçalar
    # "alaka kapısı" sorgusuyla eşleşebilsin diye başlık her parçaya eklenir.
    title = re.sub(r"[-_]+", " ", Path(name).stem).strip()
    vectors = embedder.embed([f"{title}\n{chunk}" for chunk in chunks])
    existing = store.document_by_name(name)
    if existing:
        store.delete_document(existing["id"])
    doc_id = store.add_document(
        name=name, source=source, kind=kind, size=len(text.encode("utf-8")),
        chunks=chunks, vectors=vectors, model=embedder.name,
    )
    return {"id": doc_id, "name": name, "chunks": len(chunks), "kind": kind}


def index_file(path: Path, embedder: BaseEmbedder, settings: Dict[str, Any]) -> Dict[str, Any]:
    text = read_file(path)
    return index_text(
        name=path.name, text=text, source=str(path),
        kind=path.suffix.lower().lstrip("."), embedder=embedder, settings=settings,
    )


def scan_documents_folder() -> List[Path]:
    return sorted(
        p for p in DOCS_DIR.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED
    )


def reindex_all(embedder: BaseEmbedder, settings: Dict[str, Any],
                progress: Optional[Callable[[int, int, str], None]] = None) -> Dict[str, Any]:
    """documents/ klasörünü baştan tarayıp tüm indeksi yeniden kurar."""
    files = scan_documents_folder()
    store.clear_all()
    results: List[Dict[str, Any]] = []
    errors: List[Dict[str, str]] = []
    total = len(files)
    for i, path in enumerate(files, start=1):
        if progress:
            progress(i, total, path.name)
        try:
            results.append(index_file(path, embedder, settings))
        except Exception as exc:
            errors.append({"name": path.name, "error": str(exc)})
    store.set_meta("embedder", embedder.name)
    store.set_meta("embedder_kind", embedder.kind)
    return {
        "indexed": len(results),
        "chunks": sum(r["chunks"] for r in results),
        "files": results,
        "errors": errors,
    }
